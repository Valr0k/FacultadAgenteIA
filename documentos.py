"""
Módulo de documentos: lectura de archivos, división en fragmentos y búsqueda.

Formatos soportados:
- PDF (.pdf)             -> cita por página
- Word (.docx)           -> cita por sección (título) y página aproximada
- PowerPoint (.pptx)     -> cita por diapositiva
- Texto (.txt, .md)      -> cita por rango de líneas

Cada fragmento guarda el nombre del archivo y su ubicación dentro de él,
para que el asistente pueda citar la fuente exacta y el usuario la busque a mano.
"""

import io
import re
from dataclasses import dataclass
from pathlib import Path

from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Palabras muy comunes en español que no aportan a la búsqueda
PALABRAS_VACIAS = [
    "a", "al", "algo", "ante", "con", "como", "cual", "cuando", "de", "del",
    "desde", "donde", "el", "ella", "en", "entre", "era", "es", "esa", "ese",
    "eso", "esta", "este", "esto", "fue", "ha", "hay", "la", "las", "le",
    "les", "lo", "los", "mas", "me", "mi", "muy", "no", "o", "para", "pero",
    "por", "que", "se", "ser", "si", "sin", "sobre", "son", "su",
    "sus", "tambien", "te", "tiene", "un", "una", "uno", "unos", "unas", "y",
    "ya", "yo",
]

TAMANIO_FRAGMENTO = 1200   # caracteres por fragmento
SOLAPAMIENTO = 200         # caracteres compartidos entre fragmentos vecinos

EXTENSIONES_SOPORTADAS = ["pdf", "docx", "pptx", "txt", "md"]


@dataclass
class Fragmento:
    archivo: str
    ubicacion: str   # por ejemplo: "pág. 3", "diapositiva 5", "líneas 10–40"
    texto: str

    def cita(self) -> str:
        return f"{self.archivo}, {self.ubicacion}"


@dataclass
class Unidad:
    """Una parte mínima de un documento (un párrafo, una línea, una tabla)
    con los datos necesarios para ubicarla después."""
    numero: int
    texto: str
    seccion: str | None = None
    pagina: int | None = None


# =====================================================================
# Utilidades de texto
# =====================================================================

def limpiar_texto(texto: str) -> str:
    """Quita espacios repetidos y saltos de línea innecesarios."""
    texto = texto.replace("\x00", " ")
    texto = re.sub(r"-\n(\w)", r"\1", texto)      # une palabras cortadas con guion
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def dividir_texto(texto: str) -> list[str]:
    """Divide un texto largo en partes con solapamiento,
    para no cortar una idea justo en el borde."""
    if len(texto) <= TAMANIO_FRAGMENTO:
        return [texto]
    partes = []
    inicio = 0
    while inicio < len(texto):
        fin = inicio + TAMANIO_FRAGMENTO
        # Intenta cortar al final de una oración para no partir frases
        if fin < len(texto):
            punto = texto.rfind(". ", inicio + TAMANIO_FRAGMENTO // 2, fin)
            if punto != -1:
                fin = punto + 1
        partes.append(texto[inicio:fin].strip())
        if fin >= len(texto):
            break
        inicio = fin - SOLAPAMIENTO
    return partes


def rango(desde: int, hasta: int) -> str:
    return str(desde) if desde == hasta else f"{desde}–{hasta}"


def rango_con_nombre(singular: str, plural: str, desde: int, hasta: int) -> str:
    """Ej.: 'párrafo 3' o 'párrafos 3–7'."""
    return f"{singular} {desde}" if desde == hasta else f"{plural} {desde}–{hasta}"


def agrupar_unidades(archivo: str, unidades: list[Unidad],
                     describir_ubicacion) -> list[Fragmento]:
    """Junta unidades chicas (párrafos, líneas) hasta llegar al tamaño de un
    fragmento. Corta también cuando cambia la sección, para no mezclar temas.
    `describir_ubicacion(grupo)` arma el texto de la ubicación para la cita."""
    fragmentos = []
    grupo: list[Unidad] = []
    largo = 0

    def cerrar_grupo():
        if not grupo:
            return
        texto = " ".join(u.texto for u in grupo)
        ubicacion = describir_ubicacion(grupo)
        for parte in dividir_texto(texto):
            fragmentos.append(Fragmento(archivo=archivo, ubicacion=ubicacion, texto=parte))

    for unidad in unidades:
        cambia_seccion = grupo and unidad.seccion != grupo[-1].seccion
        if grupo and (largo + len(unidad.texto) > TAMANIO_FRAGMENTO or cambia_seccion):
            cerrar_grupo()
            grupo, largo = [], 0
        grupo.append(unidad)
        largo += len(unidad.texto) + 1
    cerrar_grupo()
    return fragmentos


# =====================================================================
# Lectores por formato
# Cada uno devuelve (fragmentos, avisos). Los avisos se muestran al usuario.
# =====================================================================

def leer_pdf(nombre: str, contenido: bytes) -> tuple[list[Fragmento], list[str]]:
    """PDF: se lee página por página. La cita es el número de página del archivo."""
    lector = PdfReader(io.BytesIO(contenido))
    fragmentos = []
    paginas_vacias = 0
    for numero, pagina in enumerate(lector.pages, start=1):
        texto = limpiar_texto(pagina.extract_text() or "")
        if len(texto) < 20:
            paginas_vacias += 1
            continue
        for parte in dividir_texto(texto):
            fragmentos.append(Fragmento(archivo=nombre, ubicacion=f"pág. {numero}", texto=parte))
    avisos = []
    if paginas_vacias:
        avisos.append(f"{paginas_vacias} página(s) sin texto legible (pueden ser imágenes escaneadas).")
    return fragmentos, avisos


def es_titulo(parrafo) -> bool:
    """Detecta si un párrafo de Word usa un estilo de título."""
    try:
        estilo = (parrafo.style.name or "").lower()
    except Exception:
        return False
    return estilo.startswith(("heading", "título", "titulo", "title"))


def contar_saltos_de_pagina(elemento) -> int:
    """Word guarda marcas donde terminó cada página la última vez que se abrió
    el archivo. Contarlas permite estimar el número de página."""
    return len(elemento.xpath('.//w:lastRenderedPageBreak | .//w:br[@w:type="page"]'))


def texto_de_tabla(tabla) -> str:
    """Convierte una tabla de Word en texto, fila por fila."""
    filas = []
    for fila in tabla.rows:
        celdas = []
        for celda in fila.cells:
            texto = limpiar_texto(celda.text)
            if texto and (not celdas or celdas[-1] != texto):   # evita celdas combinadas repetidas
                celdas.append(texto)
        if celdas:
            filas.append(" | ".join(celdas))
    return "Tabla: " + " / ".join(filas) if filas else ""


def leer_docx(nombre: str, contenido: bytes) -> tuple[list[Fragmento], list[str]]:
    """Word: se recorren párrafos y tablas en orden.
    Como Word no guarda números de página fijos, se cita la sección (último
    título) y una página aproximada si el archivo tiene las marcas necesarias."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    documento = Document(io.BytesIO(contenido))
    unidades = []
    seccion_actual = None
    pagina_actual = 1
    hay_marcas_de_pagina = False
    numero_parrafo = 0

    for elemento in documento.element.body.iterchildren():
        etiqueta = elemento.tag.split("}")[-1]
        if etiqueta == "p":
            bloque = Paragraph(elemento, documento)
            texto = limpiar_texto(bloque.text)
        elif etiqueta == "tbl":
            bloque = Table(elemento, documento)
            texto = texto_de_tabla(bloque)
        else:
            continue

        pagina_inicio = pagina_actual
        saltos = contar_saltos_de_pagina(elemento)
        if saltos:
            hay_marcas_de_pagina = True
            pagina_actual += saltos

        if not texto:
            continue
        if etiqueta == "p" and es_titulo(bloque):
            seccion_actual = texto[:80]
        numero_parrafo += 1
        unidades.append(Unidad(numero=numero_parrafo, texto=texto,
                               seccion=seccion_actual, pagina=pagina_inicio))

    def describir(grupo: list[Unidad]) -> str:
        partes = []
        if hay_marcas_de_pagina:
            partes.append(f"pág. aprox. {rango(grupo[0].pagina, grupo[-1].pagina)}")
        if grupo[0].seccion:
            partes.append(f"sección «{grupo[0].seccion}»")
        partes.append(rango_con_nombre("párrafo", "párrafos", grupo[0].numero, grupo[-1].numero))
        return ", ".join(partes)

    avisos = []
    if not unidades:
        avisos.append("no tiene texto legible.")
    elif not hay_marcas_de_pagina:
        avisos.append("no trae información de páginas: se citará por sección y número de párrafo.")
    return agrupar_unidades(nombre, unidades, describir), avisos


def textos_de_formas(formas) -> list[str]:
    """Saca el texto de las formas de una diapositiva (cuadros de texto,
    tablas y grupos de formas)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    textos = []
    for forma in formas:
        if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
            textos.extend(textos_de_formas(forma.shapes))
        elif getattr(forma, "has_table", False) and forma.has_table:
            filas = [" | ".join(limpiar_texto(c.text) for c in fila.cells if c.text.strip())
                     for fila in forma.table.rows]
            textos.append("Tabla: " + " / ".join(f for f in filas if f))
        elif getattr(forma, "has_text_frame", False) and forma.has_text_frame:
            texto = limpiar_texto(forma.text_frame.text)
            if texto:
                textos.append(texto)
    return textos


def leer_pptx(nombre: str, contenido: bytes) -> tuple[list[Fragmento], list[str]]:
    """PowerPoint: se lee diapositiva por diapositiva, incluyendo las notas
    del orador. La cita es el número de diapositiva."""
    from pptx import Presentation

    presentacion = Presentation(io.BytesIO(contenido))
    fragmentos = []
    diapositivas_vacias = 0
    for numero, diapositiva in enumerate(presentacion.slides, start=1):
        textos = textos_de_formas(diapositiva.shapes)
        if diapositiva.has_notes_slide:
            notas = limpiar_texto(diapositiva.notes_slide.notes_text_frame.text or "")
            if notas:
                textos.append(f"Notas del orador: {notas}")
        texto = " ".join(textos)
        if len(texto) < 5:
            diapositivas_vacias += 1
            continue
        for parte in dividir_texto(texto):
            fragmentos.append(Fragmento(archivo=nombre, ubicacion=f"diapositiva {numero}", texto=parte))
    avisos = []
    if diapositivas_vacias:
        avisos.append(f"{diapositivas_vacias} diapositiva(s) sin texto (pueden tener solo imágenes).")
    return fragmentos, avisos


def decodificar_texto(contenido: bytes) -> str:
    """Prueba las codificaciones más comunes en archivos de texto en español."""
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("utf-8", errors="replace")


def leer_texto(nombre: str, contenido: bytes) -> tuple[list[Fragmento], list[str]]:
    """Texto plano: se cita por rango de líneas (se ven en cualquier editor)."""
    lineas = decodificar_texto(contenido).splitlines()
    unidades = [Unidad(numero=i, texto=limpiar_texto(linea))
                for i, linea in enumerate(lineas, start=1) if linea.strip()]

    def describir(grupo: list[Unidad]) -> str:
        return rango_con_nombre("línea", "líneas", grupo[0].numero, grupo[-1].numero)

    avisos = [] if unidades else ["está vacío."]
    return agrupar_unidades(nombre, unidades, describir), avisos


LECTORES = {
    "pdf": leer_pdf,
    "docx": leer_docx,
    "pptx": leer_pptx,
    "txt": leer_texto,
    "md": leer_texto,
}


def leer_documento(nombre: str, contenido: bytes) -> tuple[list[Fragmento], list[str]]:
    """Elige el lector según la extensión del archivo."""
    extension = Path(nombre).suffix.lower().lstrip(".")
    lector = LECTORES.get(extension)
    if lector is None:
        if extension in ("doc", "ppt"):
            return [], [f"el formato viejo .{extension} no se puede leer. Abrilo y "
                        f"guardalo como .{extension}x (Archivo → Guardar como)."]
        return [], [f"formato .{extension} no soportado."]
    try:
        return lector(nombre, contenido)
    except Exception as error:
        return [], [f"no se pudo leer ({error}). Si es un archivo protegido o dañado, "
                    "probá guardarlo de nuevo o exportarlo a PDF."]


# =====================================================================
# Búsqueda
# =====================================================================

class Buscador:
    """Encuentra los fragmentos más parecidos a una pregunta usando TF-IDF.
    Es gratuito, rápido y no gasta cuota de la API."""

    def __init__(self, fragmentos: list[Fragmento]):
        self.fragmentos = fragmentos
        self.vectorizador = TfidfVectorizer(
            strip_accents="unicode",
            lowercase=True,
            stop_words=PALABRAS_VACIAS,
            ngram_range=(1, 2),
            sublinear_tf=True,
        )
        self.matriz = self.vectorizador.fit_transform([f.texto for f in fragmentos])

    def buscar(self, consulta: str, cantidad: int = 8) -> list[tuple[Fragmento, float]]:
        vector = self.vectorizador.transform([consulta])
        puntajes = cosine_similarity(vector, self.matriz)[0]
        orden = puntajes.argsort()[::-1][:cantidad]
        return [(self.fragmentos[i], float(puntajes[i])) for i in orden if puntajes[i] > 0]


def armar_contexto(resultados: list[tuple[Fragmento, float]]) -> str:
    """Convierte los fragmentos encontrados en el texto que recibe el modelo."""
    if not resultados:
        return "NO SE ENCONTRARON FRAGMENTOS RELEVANTES EN LOS DOCUMENTOS."
    bloques = []
    for i, (fragmento, _) in enumerate(resultados, start=1):
        bloques.append(
            f"[Fragmento {i}] Documento: {fragmento.archivo} | Ubicación: {fragmento.ubicacion}\n"
            f"{fragmento.texto}"
        )
    return "\n\n".join(bloques)
